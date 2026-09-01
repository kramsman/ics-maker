"""Read the events spreadsheet and coerce cells into Python values.

Excel and Numbers hand back a surprising variety of types for what looks like
one column: a date cell may arrive as a ``datetime``, a ``date``, or plain text
that a human typed. The coerce_* helpers here absorb that variety so the rest
of the program only ever sees ``date``, ``time``, ``bool`` and ``str``.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from ics_maker import constants


class SheetError(Exception):
    """A structural problem with the workbook (missing sheet or column)."""


@dataclass
class RawRow:
    """One spreadsheet row, keyed by canonical field name.

    Attributes:
        row_number: 1-based row number as shown in Excel. Used in error
            messages and as the fallback output filename.
        values: Canonical field name -> raw cell value (None when empty).
    """

    row_number: int
    values: dict[str, object] = field(default_factory=dict)

    def raw(self, name: str) -> object:
        """Return the raw cell value for a canonical field name."""
        return self.values.get(name)

    def text(self, name: str) -> str:
        """Return a field as trimmed text, or '' when the cell is empty."""
        return clean_str(self.values.get(name))


def normalize_header(value: object) -> str:
    """Reduce a header cell to a comparable key.

    Lowercases and drops spaces, underscores, hyphens and periods, so that
    "Start Date", "start_date" and "START-DATE" all match.

    Args:
        value: The raw header cell value (any type openpyxl might return).

    Returns:
        The normalized key used to look up the field in constants.COLUMNS.
    """
    text = clean_str(value).lower()
    for ch in " _-.":  # Strip separators so header variants collapse together.
        text = text.replace(ch, "")
    return text


def clean_str(value: object) -> str:
    """Return a cell as trimmed text; empty string for None.

    Args:
        value: The raw cell value (any type openpyxl might return).

    Returns:
        The value as trimmed text, or "" if the cell was None.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def coerce_date(value: object, field_name: str) -> dt.date | None:
    """Convert a cell to a date, or None when the cell is empty.

    Args:
        value: The raw cell value: a datetime, a date, a date-like string, or
            None.
        field_name: Human-readable column name, used in the error message.

    Returns:
        The parsed date, or None if the cell was empty.

    Raises:
        ValueError: The cell holds something that is not a recognizable date.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = clean_str(value)
    for fmt in constants.DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(
        f"{field_name!r} is not a date I recognize: {text!r}. "
        f"Use a real Excel date cell or text like 08/31/2026."
    )


def coerce_time(value: object, field_name: str) -> dt.time | None:
    """Convert a cell to a time of day, or None when the cell is empty.

    Handles floats because Excel stores a bare time as a fraction of a day
    (0.5 is noon), and timedeltas because that is what openpyxl returns for
    duration-formatted cells.

    Args:
        value: The raw cell value: a datetime, a time, a timedelta, a
            fractional-day float, a time-like string, or None.
        field_name: Human-readable column name, used in the error message.

    Returns:
        The parsed time, or None if the cell was empty.

    Raises:
        ValueError: The cell holds something that is not a recognizable time.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, dt.datetime):
        return value.time()
    if isinstance(value, dt.time):
        return value
    if isinstance(value, dt.timedelta):
        # Wrap into a single day: a duration-formatted cell could technically
        # store more than 24 hours, which dt.time cannot represent.
        total = int(value.total_seconds()) % 86_400
        return dt.time(total // 3600, (total % 3600) // 60, total % 60)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel stores a bare time of day as the fraction of a day elapsed
        # (0.5 == noon); anything outside [0, 1) is not a time at all.
        if 0 <= float(value) < 1:
            total = round(float(value) * 86_400)
            return dt.time(total // 3600, (total % 3600) // 60, total % 60)
        raise ValueError(
            f"{field_name!r} is a number ({value}) that is not a time of day."
        )
    text = clean_str(value).upper().replace(".", "")
    for fmt in constants.TIME_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    raise ValueError(
        f"{field_name!r} is not a time I recognize: {clean_str(value)!r}. "
        f"Use a real Excel time cell or text like 19:00 or 7:00 PM."
    )


def coerce_bool(value: object, field_name: str) -> bool:
    """Convert a yes/no style cell to a bool.

    Args:
        value: The raw cell value: a bool, a yes/no-like string, or None.
        field_name: Human-readable column name, used in the error message.

    Returns:
        True or False. An empty/None cell is treated as False.

    Raises:
        ValueError: The cell holds something that is neither true nor false.
    """
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    text = clean_str(value).lower()
    if text in constants.TRUE_VALUES:
        return True
    if text in constants.FALSE_VALUES:
        return False
    raise ValueError(
        f"{field_name!r} should be yes or no, not {text!r}."
    )


def read_rows(path: Path, sheet_name: str | None = None) -> list[RawRow]:
    """Read an .xlsx file into RawRows.

    Args:
        path: Path to the workbook.
        sheet_name: Worksheet to read; the active sheet when None.

    Returns:
        One RawRow per non-empty data row, in sheet order.

    Raises:
        SheetError: The sheet, the header row, or a required column is missing.
    """
    # data_only=True so formula cells yield their last cached result rather
    # than the formula text.
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name is None:
            worksheet = workbook.active
        elif sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            available = ", ".join(workbook.sheetnames)
            raise SheetError(
                f"No sheet named {sheet_name!r} in {path.name}. Found: {available}"
            )

        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise SheetError(f"{path.name} is empty.")

    header_to_field = {
        normalize_header(header): name for name, header in constants.COLUMNS.items()
    }
    # Column index -> canonical field name.
    index_to_field: dict[int, str] = {}
    for index, cell in enumerate(rows[0]):
        field_name = header_to_field.get(normalize_header(cell))
        # Only the first column matching a given field wins, so a spreadsheet
        # with an accidental duplicate header doesn't map two columns onto
        # the same field.
        if field_name is not None and field_name not in index_to_field.values():
            index_to_field[index] = field_name

    missing = [
        constants.COLUMNS[name]
        for name in constants.REQUIRED_COLUMNS
        if name not in index_to_field.values()
    ]
    if missing:
        found = ", ".join(clean_str(c) for c in rows[0] if clean_str(c)) or "nothing"
        raise SheetError(
            f"{path.name} is missing required column(s): {', '.join(missing)}. "
            f"Row 1 must be a header row; it currently holds: {found}"
        )

    raw_rows: list[RawRow] = []
    for offset, cells in enumerate(rows[1:], start=2):
        if all(clean_str(cell) == "" for cell in cells):
            continue  # Blank spacer row, not an error.
        values = {
            name: (cells[index] if index < len(cells) else None)
            for index, name in index_to_field.items()
        }
        raw_rows.append(RawRow(row_number=offset, values=values))

    return raw_rows
