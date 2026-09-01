"""Write a starter workbook so nobody has to build the sheet from the docs."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ics_maker import constants

DATE_FORMAT = "mm/dd/yyyy"
TIME_FORMAT = "h:mm AM/PM"

# Roughly how wide each column needs to be to read comfortably in Excel.
COLUMN_WIDTHS = {
    "title": 28,
    "location": 26,
    "start_date": 12,
    "start_time": 11,
    "end_date": 12,
    "end_time": 11,
    "all_day": 9,
    "description": 34,
    "reminder": 11,
    "timezone": 20,
    "url": 26,
    "filename": 18,
}


def _example_rows(today: dt.date) -> list[dict[str, object]]:
    """Build example rows that exercise every feature worth demonstrating.

    Args:
        today: Anchor date; example rows are offset relative to this so the
            generated template always shows upcoming, not past, dates.

    Returns:
        One dict per example row, keyed by canonical field name (matching
        constants.COLUMNS). Fields not relevant to a given example are
        simply omitted rather than set to None.
    """
    return [
        {
            "title": "Board Meeting",
            "location": "Conference Room B",
            "start_date": today + dt.timedelta(days=7),
            "start_time": dt.time(19, 0),
            "end_time": dt.time(20, 30),
            "description": "Quarterly review.\nBring the printed agenda.",
            "reminder": "15m",
            "filename": "board-meeting",
        },
        {
            "title": "Company Holiday",
            "start_date": today + dt.timedelta(days=14),
            "all_day": "yes",
            "description": "Office closed.",
        },
        {
            "title": "Conference",
            "location": "Javits Center, New York",
            "start_date": today + dt.timedelta(days=30),
            "end_date": today + dt.timedelta(days=32),
            "all_day": "yes",
            "reminder": "1d",
            "url": "https://example.com/conference",
            "filename": "conference",
        },
        {
            # No Filename: this row is written as ics-<its row number>.ics.
            "title": "Call with West Coast team",
            "start_date": today + dt.timedelta(days=10),
            "start_time": dt.time(9, 0),
            "reminder": "10m",
            "timezone": "America/Los_Angeles",
            "description": "End time left blank, so this runs one hour.",
        },
    ]


def write_template(path: Path) -> Path:
    """Create a template workbook at ``path``, overwriting any existing file.

    Args:
        path: Destination .xlsx path. Parent directories are created if they
            do not already exist.

    Returns:
        The path that was written.
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Events"

    # Column order follows constants.COLUMNS so the template always matches
    # what read_rows() expects, without hardcoding the list twice.
    fields = list(constants.COLUMNS)
    for index, name in enumerate(fields, start=1):
        cell = worksheet.cell(row=1, column=index, value=constants.COLUMNS[name])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        worksheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS[name]

    for row_offset, example in enumerate(_example_rows(dt.date.today()), start=2):
        for index, name in enumerate(fields, start=1):
            value = example.get(name)
            if value is None:
                continue
            cell = worksheet.cell(row=row_offset, column=index, value=value)
            # Apply Excel's real date/time number formats (rather than
            # leaving cells as generic numbers) so they display and sort
            # correctly, and so read_rows() sees native date/time objects.
            if isinstance(value, dt.date):
                cell.number_format = DATE_FORMAT
            elif isinstance(value, dt.time):
                cell.number_format = TIME_FORMAT
            elif name == "description":
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Keep the headers visible while scrolling a long list of events.
    worksheet.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
